import asyncio
import io
import logging
import os
import sqlite3
import zipfile
from datetime import datetime, date, time, timedelta

import s3_storage
from config import DB_NAME, UPLOAD_DIR

logger = logging.getLogger(__name__)

BACKUP_HOUR = 3
REPORT_HOUR = 3
REPORT_MINUTE = 30

BACKUP_RETENTION = 30


def _seconds_until(hour: int, minute: int = 0) -> float:
    now = datetime.now()
    target = now.replace(hour=hour, minute=minute, second=0, microsecond=0)
    if target <= now:
        target += timedelta(days=1)
    return (target - now).total_seconds()


def _backup_db() -> bytes | None:
    """Safe SQLite backup using backup API (handles WAL correctly)."""
    try:
        src = sqlite3.connect(DB_NAME)
        dst = sqlite3.connect(":memory:")
        src.backup(dst)
        src.close()
        buf = io.BytesIO()
        for line in dst.iterdump():
            buf.write((line + "\n").encode("utf-8"))
        dst.close()
        return buf.getvalue()
    except Exception as e:
        logger.warning(f"DB backup via backup API failed: {e}, trying fallback")

    try:
        conn = sqlite3.connect(DB_NAME)
        conn.execute("PRAGMA wal_checkpoint(TRUNCATE)")
        conn.close()
        with open(DB_NAME, "rb") as f:
            return f.read()
    except Exception as e:
        logger.error(f"DB backup failed: {e}")
        return None


def _get_photo_bytes(filename: str) -> bytes | None:
    """Get photo from local or S3."""
    local_path = os.path.join(UPLOAD_DIR, filename)
    if os.path.isfile(local_path):
        with open(local_path, "rb") as f:
            return f.read()
    if s3_storage.is_configured():
        data = s3_storage.download(f"photos/{filename}")
        if data:
            return data
        return s3_storage.download(filename)
    return None


async def _get_months_to_report() -> list[tuple[date, date]]:
    """Determine which months need reports: current month + months with recent changes."""
    from database import AsyncSessionLocal
    from models import Transaction, AuditLog
    from sqlalchemy import select

    now = datetime.now()
    current_first = date(now.year, now.month, 1)

    # Always include current month
    months = {(current_first.year, current_first.month)}

    # Check audit log for transaction changes in last 24h
    yesterday = now - timedelta(hours=24)
    async with AsyncSessionLocal() as db:
        audit_q = (
            select(AuditLog.entity_id)
            .where(
                AuditLog.entity_type == "transaction",
                AuditLog.action.in_(["create", "update", "delete"]),
                AuditLog.created_at >= yesterday,
            )
        )
        audit_res = await db.execute(audit_q)
        tx_ids = [r[0] for r in audit_res.all() if r[0]]

        if tx_ids:
            tx_q = select(Transaction.date).where(Transaction.id.in_(tx_ids))
            tx_res = await db.execute(tx_q)
            for (tx_date,) in tx_res.all():
                if tx_date:
                    months.add((tx_date.year, tx_date.month))

    # Convert to (first_day, last_day) ranges
    result = []
    for year, month in sorted(months):
        first = date(year, month, 1)
        if month == 12:
            last = date(year, 12, 31)
        else:
            last = date(year, month + 1, 1) - timedelta(days=1)
        result.append((first, last))

    return result


async def _generate_month_report(d_from: date, d_to: date) -> tuple[bytes | None, list[str]]:
    """Generate Excel report for a specific month. Returns (excel_bytes, photo_filenames)."""
    from compliance import _build_summary_sheet, _build_operations_sheet, _build_legal_sheet
    from openpyxl import Workbook
    from database import AsyncSessionLocal
    from models import Transaction, Collection, User
    from sqlalchemy import select

    dt_from = datetime.combine(d_from, time.min)
    dt_to = datetime.combine(d_to, time.max)

    async with AsyncSessionLocal() as db:
        q = (
            select(Transaction)
            .where(Transaction.date >= dt_from, Transaction.date <= dt_to)
            .order_by(Transaction.date)
        )
        transactions = (await db.execute(q)).scalars().all()
        if not transactions:
            return None, []

        coll_ids = set(t.collection_id for t in transactions if t.collection_id)
        user_ids = set(t.created_by for t in transactions if t.created_by)

        collections_map = {}
        if coll_ids:
            res = await db.execute(select(Collection).where(Collection.id.in_(coll_ids)))
            for c in res.scalars().all():
                collections_map[c.id] = c.name

        users_map = {}
        if user_ids:
            res = await db.execute(select(User).where(User.id.in_(user_ids)))
            for u in res.scalars().all():
                users_map[u.id] = u.name

    photo_files = []
    for t in transactions:
        if t.photo_path:
            for p in t.photo_path.split(","):
                p = p.strip()
                if p:
                    photo_files.append(p)

    wb = Workbook()
    _build_summary_sheet(wb.active, transactions, d_from, d_to, collections_map, users_map)
    _build_operations_sheet(wb.create_sheet(), transactions, collections_map, users_map)
    _build_legal_sheet(wb.create_sheet())

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), photo_files


async def _upload_month_report(d_from: date, excel_data: bytes, photo_files: list[str]):
    """Upload Excel + ZIP for a specific month to S3."""
    month_str = d_from.strftime("%Y_%m")
    excel_ct = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    # 1. Standalone Excel
    excel_key = f"reports/{month_str}/financial_report_{month_str}.xlsx"
    if s3_storage.upload(excel_data, excel_key, excel_ct):
        logger.info(f"Excel uploaded: {excel_key} ({len(excel_data)} bytes)")
    else:
        logger.error(f"Excel upload failed: {excel_key}")

    # 2. ZIP: Excel + receipts
    zip_buf = io.BytesIO()
    with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(f"financial_report_{month_str}.xlsx", excel_data)
        if photo_files:
            added = 0
            for filename in photo_files:
                photo_data = await asyncio.to_thread(_get_photo_bytes, filename)
                if photo_data:
                    zf.writestr(f"receipts/{filename}", photo_data)
                    added += 1
            logger.info(f"  {added}/{len(photo_files)} photos added")

    zip_data = zip_buf.getvalue()
    zip_key = f"reports/{month_str}/financial_report_{month_str}.zip"
    if s3_storage.upload(zip_data, zip_key, "application/zip"):
        logger.info(f"Archive uploaded: {zip_key} ({len(zip_data)} bytes)")
    else:
        logger.error(f"Archive upload failed: {zip_key}")


def _cleanup_old_keys(prefix: str, keep: int):
    """Remove old S3 objects, keep only the last N."""
    try:
        client = s3_storage._get_client()
        if not client:
            return
        response = client.list_objects_v2(
            Bucket=s3_storage.S3_BUCKET, Prefix=prefix
        )
        objects = response.get("Contents", [])
        if len(objects) <= keep:
            return
        objects.sort(key=lambda o: o["Key"])
        to_delete = objects[: len(objects) - keep]
        for obj in to_delete:
            s3_storage.delete(obj["Key"])
            logger.info(f"Cleaned up: {obj['Key']}")
    except Exception as e:
        logger.error(f"Cleanup failed for {prefix}: {e}")


async def _run_backup():
    if not s3_storage.is_configured():
        return
    logger.info("Starting scheduled DB backup...")
    data = await asyncio.to_thread(_backup_db)
    if not data:
        logger.error("DB backup: no data")
        return
    key = f"backups/rodcom_{datetime.now().strftime('%Y%m%d_%H%M')}.db"
    if s3_storage.upload(data, key, "application/x-sqlite3"):
        logger.info(f"DB backup uploaded: {key} ({len(data)} bytes)")
        _cleanup_old_keys("backups/", BACKUP_RETENTION)
    else:
        logger.error("DB backup upload failed")


async def _run_reports():
    if not s3_storage.is_configured():
        return

    try:
        months = await _get_months_to_report()
    except Exception as e:
        logger.error(f"Failed to determine report months: {e}")
        return

    logger.info(f"Generating reports for {len(months)} month(s): {[m[0].strftime('%Y-%m') for m in months]}")

    for d_from, d_to in months:
        try:
            excel_data, photo_files = await _generate_month_report(d_from, d_to)
            if not excel_data:
                logger.info(f"  {d_from.strftime('%Y-%m')}: no transactions, skipping")
                continue
            await _upload_month_report(d_from, excel_data, photo_files)
        except Exception as e:
            logger.error(f"  {d_from.strftime('%Y-%m')}: failed — {e}")


async def start_scheduler():
    """Daily: DB backup at 03:00, reports at 03:30 (server timezone)."""
    if not s3_storage.is_configured():
        logger.info("S3 not configured, scheduler disabled")
        return

    logger.info("Scheduler started (backup 03:00, reports 03:30)")

    while True:
        try:
            wait = _seconds_until(BACKUP_HOUR, 0)
            logger.info(f"Next backup in {wait / 3600:.1f}h")
            await asyncio.sleep(wait)

            await _run_backup()
            await asyncio.sleep(REPORT_MINUTE * 60)
            await _run_reports()

            await asyncio.sleep(60)
        except asyncio.CancelledError:
            logger.info("Scheduler stopped")
            break
        except Exception as e:
            logger.error(f"Scheduler error: {e}")
            await asyncio.sleep(300)
